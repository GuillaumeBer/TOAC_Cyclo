"""Export mis en forme (XLSX) des sorties de la base, prévu pour être converti en Google Sheet.

Quatre onglets :
  - « Récapitulatif » : toutes les sorties, une ligne chacune, pour une vue d'ensemble (le titre ouvre la page de l'événement) ;
  - « Sorties »  : le détail de chaque sortie, circuits cliquables (un lien OpenRunner par circuit), liens plan et événement ;
  - « Circuits » : une ligne par circuit (distance, dénivelé, tracé cliquable, remarque) ;
  - « Légende »  : lecture des couleurs et limites des données.

Les URL ne s'affichent jamais en clair : le texte visible est « Route 42 km · D+ 350 m », « Voir le plan », « Voir la page »...
"""

import re
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from periode import JOURS

DISCIPLINES = {"route": "Route", "gravel": "Gravel", "vtt": "VTT"}
MAX_CIRCUIT_COLUMNS = 6

NAVY = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BANDS = ("FFFFFF", "F3F6FB")  # alternance par jour, pour repérer les sorties d'une même date
GREEN, ORANGE, RED = "C6EFCE", "FFEB9C", "FFC7CE"
FEDERATION_FILLS = {"FFVélo": "DDEBF7", "FSGT": "FCE4D6", "UFOLEP": "E2EFDA", "FFC": "EDEDED"}


def _num(value: float | None) -> str:
    return f"{value:g}".replace(".", ",") if value is not None else ""


def _label(circuit: dict, edition: str) -> str:
    """Texte visible d'un circuit : « Route 42 km · D+ 350 m », avec un signe si le tracé est titré d'une autre année."""
    parts = [DISCIPLINES.get(circuit.get("discipline"), circuit.get("discipline") or "Circuit")]
    if circuit.get("distance_km"):
        parts[0] += f" {_num(circuit['distance_km'])} km"
    text = parts[0] + (f" · D+ {_num(circuit['denivele_m'])} m" if circuit.get("denivele_m") else "")
    year = circuit.get("trace_annee")
    return text + (f"  ⚠ titré {year}" if year and year != edition else "")


def _set(ws, row: int, col: int, value, *, link: str | None = None, bold: bool = False, fill: str | None = None,
         wrap: bool = True, center: bool = False, number_format: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = LINK_FONT if link else Font(name="Arial", size=10, bold=bold)
    if link:
        cell.hyperlink = link
    cell.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal="center" if center else "left")
    cell.border = BORDER
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if number_format:
        cell.number_format = number_format


def _header(ws, headers: list[tuple[str, int]]) -> None:
    for col, (name, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _fed(s: dict) -> str:
    federation = s.get("federation")
    return "Autre" if federation == "autre" else federation or "Non précisée"


def _http(url: str | None) -> str | None:
    return url if url and str(url).startswith("http") else None


def _sorties_sheet(ws, sorties: list[dict], maps: Callable[[dict], str]) -> None:
    headers = [("Date", 12), ("Jour", 11), ("Titre", 34), ("Disciplines", 14), ("Fédération", 13), ("Commune", 18),
               ("RDV", 8), ("Départ", 8), ("Lieu de RDV", 34), ("Plan", 13)]
    headers += [(f"Circuit {i}", 26) for i in range(1, MAX_CIRCUIT_COLUMNS + 1)]
    headers += [("Autres circuits", 20), ("Tarif", 42), ("Page de l'événement", 18), ("Organisateur", 28), ("Contact", 34),
                ("Fiabilité", 17), ("Statut", 11), ("Dernière vérification", 17)]
    _header(ws, headers)

    band, previous_day = 0, None
    for row, s in enumerate(sorties, 2):
        if s["date"] != previous_day:
            band, previous_day = 1 - band, s["date"]
        fill = BANDS[band]
        edition = s["date"][:4]
        day = date.fromisoformat(s["date"])
        circuits = sorted((c for c in s.get("circuits") or [] if c.get("distance_km") or c.get("lien_trace")),
                          key=lambda c: (c.get("discipline") or "", c.get("distance_km") or 0))
        unnamed = len(s.get("circuits") or []) - len(circuits)
        verified = s.get("fiabilite") == "confirmee"
        address = re.sub(r"\s*\([^)]*\)", "", s.get("adresse_depart") or "").strip(" ,")

        values = [
            (day, {"number_format": "dd/mm/yyyy", "center": True}),
            (JOURS[day.weekday()].capitalize(), {}),
            (s.get("titre") or "", {"bold": True}),
            (" · ".join(DISCIPLINES.get(d, d) for d in s.get("disciplines") or []), {}),
            (_fed(s), {"fill": FEDERATION_FILLS.get(s.get("federation") or "")}),
            (s.get("lieu_depart") or "", {}),
            (s.get("heure_rdv") or "", {"center": True}),
            (s.get("heure_depart") or "", {"center": True}),
            (address or s.get("lieu_depart") or "", {}),
        ]
        for col, (value, opts) in enumerate(values, 1):
            _set(ws, row, col, value, **{"fill": opts.pop("fill", None) or fill, **opts})
        has_place = s.get("adresse_depart") or s.get("lieu_depart")
        _set(ws, row, 10, "Voir le plan" if has_place else "", link=maps(s) if has_place else None, fill=fill)

        for i in range(MAX_CIRCUIT_COLUMNS):  # un circuit par colonne, cliquable quand le tracé est connu
            col = 11 + i
            if i < len(circuits):
                c = circuits[i]
                _set(ws, row, col, _label(c, edition), link=_http(c.get("lien_trace")), fill=fill)
            else:
                _set(ws, row, col, "", fill=fill)
        extra = max(len(circuits) - MAX_CIRCUIT_COLUMNS, 0)
        notes = ([f"+ {extra} circuit(s) : voir l'onglet Circuits"] if extra else []) + (
            [f"{unnamed} circuit(s) annoncé(s), distances non publiées"] if unnamed else [])
        col = 11 + MAX_CIRCUIT_COLUMNS
        _set(ws, row, col, "\n".join(notes), fill=fill)
        _set(ws, row, col + 1, s.get("tarif_inscription") or "", fill=fill)
        _set(ws, row, col + 2, "Voir la page" if _http(s.get("url")) else "", link=_http(s.get("url")), fill=fill)
        _set(ws, row, col + 3, s.get("organisateur") or "", fill=fill)
        _set(ws, row, col + 4, s.get("contact") or "", fill=fill)
        _set(ws, row, col + 5, "Vérifiée chez l'organisateur" if verified else "À vérifier", fill=GREEN if verified else ORANGE, center=True)
        _set(ws, row, col + 6, "Annulée" if s.get("statut") == "annulee" else "Maintenue",
             fill=RED if s.get("statut") == "annulee" else fill, center=True)
        _set(ws, row, col + 7, (s.get("derniere_vue") or "")[:16].replace("T", " "), fill=fill, center=True)


def _distances(s: dict) -> str:
    """Résumé des distances par discipline : « Route 31 / 72 / 98 km · VTT 10 / 25 / 32 km »."""
    by_discipline: dict[str, list[float]] = {}
    for c in s.get("circuits") or []:
        if c.get("distance_km"):
            by_discipline.setdefault(DISCIPLINES.get(c.get("discipline"), c.get("discipline") or "Circuit"), []).append(c["distance_km"])
    return " · ".join(f"{name} {' / '.join(_num(d) for d in sorted(values))} km" for name, values in by_discipline.items()) or "Non publiées"


def _recap_sheet(ws, sorties: list[dict]) -> None:
    """Vue d'ensemble : toutes les sorties sur une ligne chacune, pour repérer l'essentiel d'un coup d'œil."""
    _header(ws, [("Date", 12), ("Jour", 11), ("Horaire", 15), ("Sortie", 38), ("Commune", 18), ("Disciplines", 14),
                 ("Fédération", 13), ("Distances", 34), ("Tracés", 10), ("Fiabilité", 17)])
    ws.freeze_panes = "E2"  # date, jour, horaire et titre restent visibles au défilement

    band, previous_day = 0, None
    for row, s in enumerate(sorties, 2):
        if s["date"] != previous_day:
            band, previous_day = 1 - band, s["date"]
        fill = BANDS[band]
        day = date.fromisoformat(s["date"])
        circuits = s.get("circuits") or []
        with_links = sum(1 for c in circuits if c.get("lien_trace"))
        verified = s.get("fiabilite") == "confirmee"
        rdv, depart = s.get("heure_rdv"), s.get("heure_depart")
        hours = f"RDV {rdv} · départ {depart}" if rdv and depart and rdv != depart else (rdv or depart or "À confirmer")

        _set(ws, row, 1, day, number_format="dd/mm/yyyy", center=True, fill=fill)
        _set(ws, row, 2, JOURS[day.weekday()].capitalize(), fill=fill)
        _set(ws, row, 3, hours, center=True, fill=fill)
        page = _http(s.get("url"))
        if page:  # le titre ouvre la page de l'événement
            _set(ws, row, 4, s.get("titre") or "", link=page, fill=fill)
            ws.cell(row=row, column=4).font = Font(name="Arial", size=10, bold=True, color="0563C1", underline="single")
        else:
            _set(ws, row, 4, s.get("titre") or "", bold=True, fill=fill)
        _set(ws, row, 5, s.get("lieu_depart") or "", fill=fill)
        _set(ws, row, 6, " · ".join(DISCIPLINES.get(d, d) for d in s.get("disciplines") or []), fill=fill)
        _set(ws, row, 7, _fed(s), fill=FEDERATION_FILLS.get(s.get("federation") or "") or fill)
        _set(ws, row, 8, _distances(s), fill=fill)
        _set(ws, row, 9, f"{with_links}/{len(circuits)}" if circuits else "—", center=True, fill=fill)
        _set(ws, row, 10, "Vérifiée chez l'organisateur" if verified else "À vérifier", fill=GREEN if verified else ORANGE, center=True)


def _circuits_sheet(ws, sorties: list[dict]) -> None:
    _header(ws, [("Date", 12), ("Sortie", 36), ("Commune", 18), ("Discipline", 12), ("Distance (km)", 13),
                 ("Dénivelé (m)", 13), ("Tracé", 16), ("Remarque", 34)])
    ws.freeze_panes = "C2"
    row = 2
    for s in sorties:
        edition = s["date"][:4]
        for c in sorted(s.get("circuits") or [], key=lambda c: (c.get("discipline") or "", c.get("distance_km") or 0)):
            year = c.get("trace_annee")
            remark = (f"⚠ Tracé titré {year} (édition {edition}) : à confirmer" if year and year != edition
                      else "" if c.get("distance_km") else "Distance non publiée")
            link = _http(c.get("lien_trace"))
            _set(ws, row, 1, date.fromisoformat(s["date"]), number_format="dd/mm/yyyy", center=True)
            _set(ws, row, 2, s.get("titre") or "", bold=True)
            _set(ws, row, 3, s.get("lieu_depart") or "")
            _set(ws, row, 4, DISCIPLINES.get(c.get("discipline"), c.get("discipline") or ""))
            _set(ws, row, 5, c.get("distance_km"), number_format="0.0", center=True)
            _set(ws, row, 6, c.get("denivele_m"), number_format="0", center=True)
            _set(ws, row, 7, "Ouvrir le tracé" if link else "", link=link)
            _set(ws, row, 8, remark, fill=ORANGE if remark.startswith("⚠") else None)
            row += 1
    ws.auto_filter.ref = f"A1:H{max(row - 1, 1)}"


def _legend_sheet(ws, sorties: list[dict]) -> None:
    lines = [
        ("Sorties cyclo autour de Toulouse", True),
        (f"Généré le {datetime.now():%d/%m/%Y à %H:%M} — {len(sorties)} sortie(s). Données recueillies sur internet par un agent : à recouper avant de partir.", False),
        ("", False),
        ("Fiabilité", True),
        ("Vérifiée chez l'organisateur : l'édition de l'année a été retrouvée sur le site de l'organisateur.", False),
        ("À vérifier : information issue d'un calendrier ou d'un agrégateur, non confirmée chez l'organisateur.", False),
        ("", False),
        ("Tracés OpenRunner", True),
        ("Chaque circuit est un lien cliquable. « ⚠ titré 2025 » signale un tracé dont le titre porte une autre année que l'édition :", False),
        ("le club le réutilise peut-être, mais ce n'est pas confirmé.", False),
        ("", False),
        ("Fédérations", True),
        ("Indiquées seulement d'après un indice explicite (logo, calendrier fédéral, club affilié). « Non précisée » : aucun indice trouvé.", False),
        ("", False),
        ("Contact", True),
        ("Coordonnées publiées sur les pages des événements, dont des personnes citées nommément : à retirer avant de partager ce fichier.", False),
    ]
    ws.column_dimensions["A"].width = 130
    for row, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name="Arial", size=12 if row == 1 else 10, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def ecrire_xlsx(chemin: Path, sorties: list[dict], maps: Callable[[dict], str]) -> Path:
    """Écrit le classeur mis en forme. `maps` construit le lien Google Maps (même fonction que pour les messages)."""
    ordered = sorted(sorties, key=lambda s: (s["date"], s.get("heure_depart") or s.get("heure_rdv") or "99:99", s.get("titre") or ""))
    wb = Workbook()
    _recap_sheet(wb.active, ordered)
    wb.active.title = "Récapitulatif"
    _sorties_sheet(wb.create_sheet("Sorties"), ordered, maps)
    _circuits_sheet(wb.create_sheet("Circuits"), ordered)
    _legend_sheet(wb.create_sheet("Légende"), ordered)
    chemin.parent.mkdir(exist_ok=True)
    wb.save(chemin)
    return chemin
